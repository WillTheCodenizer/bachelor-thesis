from __future__ import annotations

import re
import tempfile
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SCRIPT_PATH = ROOT / "sprechskript_bachelorarbeit.md"
OUTPUT_PATH = ROOT / "bachelor_thesis_presentation.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(24, 35, 49)
NAVY = RGBColor(28, 55, 80)
TEAL = RGBColor(24, 137, 133)
CORAL = RGBColor(221, 91, 74)
GOLD = RGBColor(221, 163, 58)
SKY = RGBColor(92, 155, 189)
PAPER = RGBColor(247, 245, 239)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(91, 103, 113)
PALE_TEAL = RGBColor(224, 239, 235)
PALE_CORAL = RGBColor(247, 228, 221)
PALE_BLUE = RGBColor(226, 237, 244)
LINE = RGBColor(210, 215, 215)

TITLE_FONT = "Aptos Display"
BODY_FONT = "Aptos"


def parse_speaker_notes(markdown: str) -> dict[int, str]:
    notes: dict[int, str] = {}
    pattern = re.compile(
        r"^## Slide (\d+).*?^### Speaker notes\s*\n(.*?)(?=\n---|\Z)",
        re.MULTILINE | re.DOTALL,
    )
    for match in pattern.finditer(markdown):
        notes[int(match.group(1))] = match.group(2).strip()
    return notes


def set_background(slide, color: RGBColor = PAPER) -> None:
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = color


def add_rect(slide, x, y, width, height, color, radius=False, line_color=None):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if radius
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, x, y, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line_color or color
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def set_text(
    shape,
    text: str,
    size: float,
    color: RGBColor = INK,
    bold: bool = False,
    font: str = BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margins=(0.04, 0.04, 0.04, 0.04),
):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margins[0])
    frame.margin_right = Inches(margins[1])
    frame.margin_top = Inches(margins[2])
    frame.margin_bottom = Inches(margins[3])
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_text(slide, text, x, y, width, height, size, **kwargs):
    box = slide.shapes.add_textbox(x, y, width, height)
    return set_text(box, text, size, **kwargs)


def add_rich_text(slide, segments, x, y, width, height, size=22, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    for text, color, bold in segments:
        run = paragraph.add_run()
        run.text = text
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_title(slide, number: int, title: str, kicker: str | None = None) -> None:
    add_rect(slide, Inches(0), Inches(0), Inches(0.13), SLIDE_H, TEAL)
    add_text(
        slide,
        f"{number:02d}",
        Inches(0.55),
        Inches(0.35),
        Inches(0.62),
        Inches(0.35),
        13,
        color=CORAL,
        bold=True,
    )
    if kicker:
        add_text(
            slide,
            kicker.upper(),
            Inches(1.25),
            Inches(0.35),
            Inches(3.2),
            Inches(0.28),
            10,
            color=MUTED,
            bold=True,
        )
    add_text(
        slide,
        title,
        Inches(0.55),
        Inches(0.74),
        Inches(11.9),
        Inches(0.62),
        28,
        color=NAVY,
        bold=True,
        font=TITLE_FONT,
    )
    add_rect(slide, Inches(0.55), Inches(1.47), Inches(12.0), Inches(0.015), LINE)


def add_footer(slide, number: int) -> None:
    add_text(
        slide,
        "Will Ruppert | Bachelor Thesis",
        Inches(0.58),
        Inches(7.18),
        Inches(3.4),
        Inches(0.18),
        8.5,
        color=MUTED,
    )
    add_text(
        slide,
        str(number),
        Inches(12.1),
        Inches(7.15),
        Inches(0.55),
        Inches(0.22),
        9,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_bullets(slide, bullets, x, y, width, height, size=20, color=INK):
    box = slide.shapes.add_textbox(x, y, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    for index, item in enumerate(bullets):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level
        paragraph.font.name = BODY_FONT
        paragraph.font.size = Pt(size - level * 2)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(12 if level == 0 else 7)
        paragraph.line_spacing = 1.08
        paragraph.bullet = True
    return box


def render_pdf(pdf_path: Path, output_dir: Path) -> Path:
    target = output_dir / f"{pdf_path.stem}.png"
    document = pymupdf.open(pdf_path)
    page = document[0]
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2.6, 2.6), alpha=False)
    pixmap.save(target)
    document.close()
    return target


def image_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as image:
        return image.size


def add_image_contain(slide, path: Path, x, y, width, height, border=False):
    pixel_w, pixel_h = image_size(path)
    image_ratio = pixel_w / pixel_h
    box_ratio = width / height
    if image_ratio > box_ratio:
        draw_w = width
        draw_h = width / image_ratio
    else:
        draw_h = height
        draw_w = height * image_ratio
    draw_x = x + (width - draw_w) / 2
    draw_y = y + (height - draw_h) / 2
    picture = slide.shapes.add_picture(str(path), draw_x, draw_y, draw_w, draw_h)
    if border:
        picture.line.color.rgb = LINE
        picture.line.width = Pt(0.8)
    return picture


def add_image_cover(slide, path: Path, x, y, width, height):
    pixel_w, pixel_h = image_size(path)
    image_ratio = pixel_w / pixel_h
    box_ratio = width / height
    picture = slide.shapes.add_picture(str(path), x, y, width, height)
    if image_ratio > box_ratio:
        shown_ratio = box_ratio / image_ratio
        crop = (1 - shown_ratio) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        shown_ratio = image_ratio / box_ratio
        crop = (1 - shown_ratio) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_metric(slide, value, label, x, y, width, color):
    add_rect(slide, x, y, width, Inches(1.05), WHITE, radius=True, line_color=LINE)
    add_text(
        slide,
        value,
        x + Inches(0.18),
        y + Inches(0.14),
        width - Inches(0.36),
        Inches(0.42),
        24,
        color=color,
        bold=True,
        font=TITLE_FONT,
    )
    add_text(
        slide,
        label,
        x + Inches(0.18),
        y + Inches(0.62),
        width - Inches(0.36),
        Inches(0.24),
        10,
        color=MUTED,
        bold=True,
    )


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def build_presentation() -> None:
    notes = parse_speaker_notes(SCRIPT_PATH.read_text(encoding="utf-8"))
    if len(notes) != 13:
        raise RuntimeError(f"Expected 13 speaker-note sections, found {len(notes)}")

    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    blank = presentation.slide_layouts[6]

    with tempfile.TemporaryDirectory() as temp_name:
        temp_dir = Path(temp_name)
        figures = {
            "sky": ROOT / "programming/frb-cosmology/plots/frb/frb_sky_distribution_petroff.png",
            "bias": render_pdf(ROOT / "programming/frb-cosmology/plots/frb/FRB_bias_bz.pdf", temp_dir),
            "frb_nz": render_pdf(ROOT / "programming/frb-cosmology/plots/frb/FRB_nz_shallow_deep.pdf", temp_dir),
            "gal_nz": render_pdf(ROOT / "programming/frb-cosmology/plots/galaxy/Galaxy_nz_bins.pdf", temp_dir),
            "pk": render_pdf(ROOT / "programming/frb-cosmology/plots/power_spectrum/Pk_nonlinear.pdf", temp_dir),
            "shot": render_pdf(ROOT / "programming/frb-cosmology/plots/frb/FRB_magnetar_Cell_shallow_shotnoise.pdf", temp_dir),
            "cross": render_pdf(ROOT / "programming/frb-cosmology/plots/frb_x_galaxy/comparisons/survey_magnetar_all_bins.pdf", temp_dir),
            "fisher": render_pdf(ROOT / "programming/frb-cosmology/plots/fisher/fisher_comparison_2x2.pdf", temp_dir),
            "lsst": render_pdf(ROOT / "programming/frb-cosmology/plots/fisher_lsst/fisher_kids_vs_lsst_2x2.pdf", temp_dir),
        }

        # Slide 1
        slide = presentation.slides.add_slide(blank)
        set_background(slide, NAVY)
        add_image_cover(slide, figures["sky"], Inches(6.35), Inches(0), Inches(6.98), SLIDE_H)
        add_rect(slide, Inches(6.25), Inches(0), Inches(0.12), SLIDE_H, CORAL)
        add_text(slide, "BACHELOR THESIS", Inches(0.72), Inches(0.68), Inches(4.6), Inches(0.3), 12, color=GOLD, bold=True)
        add_text(
            slide,
            "Probing Fast Radio\nBurst Environments",
            Inches(0.72),
            Inches(1.35),
            Inches(5.0),
            Inches(1.75),
            33,
            color=WHITE,
            bold=True,
            font=TITLE_FONT,
        )
        add_text(
            slide,
            "Cross-correlating deep and shallow FRB surveys with KiDS DR5",
            Inches(0.75),
            Inches(3.42),
            Inches(4.9),
            Inches(0.9),
            17,
            color=RGBColor(215, 226, 232),
        )
        add_rect(slide, Inches(0.75), Inches(4.72), Inches(1.0), Inches(0.06), TEAL)
        add_text(slide, "WILL RUPPERT", Inches(0.75), Inches(5.1), Inches(3.2), Inches(0.35), 14, color=WHITE, bold=True)
        add_text(slide, "TU Dortmund University | 2026", Inches(0.75), Inches(5.55), Inches(3.7), Inches(0.3), 11, color=RGBColor(190, 205, 213))
        add_text(
            slide,
            "What can spatial clustering reveal about FRB progenitors?",
            Inches(0.75),
            Inches(6.35),
            Inches(5.0),
            Inches(0.52),
            14,
            color=GOLD,
            bold=True,
        )
        add_notes(slide, notes[1])

        # Slide 2
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 2, "Why Fast Radio Bursts?", "Motivation")
        add_image_contain(slide, figures["sky"], Inches(5.75), Inches(1.7), Inches(6.9), Inches(4.6))
        add_metric(slide, "< 1 ms", "TYPICAL DURATION", Inches(0.65), Inches(1.85), Inches(2.1), CORAL)
        add_metric(slide, "4,539", "BURSTS IN CHIME/FRB 2", Inches(2.95), Inches(1.85), Inches(2.25), TEAL)
        add_bullets(
            slide,
            [
                "Extreme radio luminosity",
                "Dispersion implies cosmological distances",
                "3,641 distinct sources",
                "Astrophysical origin remains uncertain",
            ],
            Inches(0.72),
            Inches(3.25),
            Inches(4.7),
            Inches(3.0),
            size=18,
        )
        add_footer(slide, 2)
        add_notes(slide, notes[2])

        # Slide 3
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 3, "Two Possible Progenitor Models", "Physical picture")
        add_rect(slide, Inches(0.68), Inches(1.82), Inches(3.18), Inches(2.05), PALE_CORAL, radius=True)
        add_text(slide, "YOUNG MAGNETARS", Inches(0.95), Inches(2.05), Inches(2.7), Inches(0.35), 16, color=CORAL, bold=True)
        add_text(slide, "Active star formation\nShort delay time\nb₀ = 1.0  |  δ = 0.8", Inches(0.95), Inches(2.57), Inches(2.6), Inches(1.05), 17, color=INK)
        add_rect(slide, Inches(0.68), Inches(4.12), Inches(3.18), Inches(2.05), PALE_BLUE, radius=True)
        add_text(slide, "DELAYED CHANNELS", Inches(0.95), Inches(4.35), Inches(2.7), Inches(0.35), 16, color=SKY, bold=True)
        add_text(slide, "Older populations\nLong delay time\nb₀ = 1.5  |  δ = 0.2", Inches(0.95), Inches(4.87), Inches(2.6), Inches(1.05), 17, color=INK)
        add_image_contain(slide, figures["bias"], Inches(4.2), Inches(1.65), Inches(8.45), Inches(4.9))
        add_rich_text(slide, [("Bias model:  ", MUTED, False), ("b(z) = b₀(1+z)ᵟ", NAVY, True)], Inches(7.05), Inches(6.25), Inches(3.5), Inches(0.4), size=18, align=PP_ALIGN.CENTER)
        add_footer(slide, 3)
        add_notes(slide, notes[3])

        # Slide 4
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 4, "The Basic Idea of Cross-Correlation", "Method")
        add_text(slide, "SAME SKY", Inches(0.72), Inches(1.85), Inches(2.1), Inches(0.3), 12, color=CORAL, bold=True)
        add_text(slide, "FRBs and galaxies trace the same underlying matter field.", Inches(0.72), Inches(2.25), Inches(3.2), Inches(1.0), 21, color=NAVY, bold=True)
        add_text(slide, "A cross-signal appears only where their redshift distributions overlap.", Inches(0.72), Inches(3.55), Inches(3.25), Inches(1.1), 17, color=MUTED)
        panel_x, panel_y = Inches(4.45), Inches(1.8)
        add_rect(slide, panel_x, panel_y, Inches(8.05), Inches(4.55), WHITE, radius=True, line_color=LINE)
        for offset, shade in [(0.0, PALE_BLUE), (1.05, PALE_TEAL), (2.1, RGBColor(240, 234, 214)), (3.15, PALE_CORAL)]:
            add_rect(slide, panel_x + Inches(0.35 + offset), panel_y + Inches(0.45), Inches(0.74), Inches(3.65), shade, radius=True)
        add_text(slide, "redshift →", panel_x + Inches(0.35), panel_y + Inches(4.08), Inches(3.8), Inches(0.3), 11, color=MUTED)
        points = [
            (1.05, 1.15, TEAL, "FRB"), (2.0, 2.7, TEAL, ""), (3.25, 1.7, TEAL, ""),
            (4.25, 2.85, TEAL, ""), (5.2, 1.3, TEAL, ""), (6.65, 2.35, TEAL, ""),
            (1.45, 2.05, CORAL, "Galaxy"), (2.65, 1.05, CORAL, ""), (3.7, 2.65, CORAL, ""),
            (4.85, 1.95, CORAL, ""), (5.9, 2.95, CORAL, ""), (6.95, 1.2, CORAL, ""),
        ]
        for px, py, color, label in points:
            circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, panel_x + Inches(px), panel_y + Inches(py), Inches(0.19), Inches(0.19))
            circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.line.color.rgb = WHITE
            if label:
                add_text(slide, label, panel_x + Inches(px + 0.25), panel_y + Inches(py - 0.05), Inches(0.8), Inches(0.25), 10, color=color, bold=True)
        add_rich_text(slide, [("tomography  ", TEAL, True), ("separates amplitude b₀ from evolution δ", INK, False)], Inches(4.7), Inches(6.52), Inches(7.4), Inches(0.38), size=16, align=PP_ALIGN.CENTER)
        add_footer(slide, 4)
        add_notes(slide, notes[4])

        # Slide 5
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 5, "Data and Survey Scenarios", "Inputs")
        add_image_contain(slide, figures["frb_nz"], Inches(0.55), Inches(1.7), Inches(6.15), Inches(3.95))
        add_image_contain(slide, figures["gal_nz"], Inches(6.68), Inches(1.7), Inches(6.05), Inches(3.95))
        add_metric(slide, "5,000", "SHALLOW FRBs | α = 3.5", Inches(0.72), Inches(5.78), Inches(2.65), CORAL)
        add_metric(slide, "50,000", "DEEP FRBs | α = 2.0", Inches(3.58), Inches(5.78), Inches(2.65), TEAL)
        add_metric(slide, "1,347 deg²", "KiDS DR5 | 6 BINS", Inches(7.3), Inches(5.78), Inches(2.65), SKY)
        add_text(slide, "Radial overlap determines the available cross-signal.", Inches(10.1), Inches(5.92), Inches(2.25), Inches(0.75), 14, color=NAVY, bold=True)
        add_footer(slide, 5)
        add_notes(slide, notes[5])

        # Slide 6
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 6, "Numerical Pipeline", "Computation")
        steps = [
            ("01", "Planck 2018", "ΛCDM cosmology", NAVY),
            ("02", "Matter field", "non-linear P(k,z)", TEAL),
            ("03", "Radial kernels", "selection × bias", CORAL),
            ("04", "Projection", "Limber → Cℓ", GOLD),
            ("05", "Forecast", "noise + Fisher", SKY),
        ]
        start_x = 0.62
        for index, (num, heading, detail, color) in enumerate(steps):
            x = Inches(start_x + index * 2.52)
            add_rect(slide, x, Inches(1.85), Inches(2.1), Inches(1.4), WHITE, radius=True, line_color=LINE)
            add_text(slide, num, x + Inches(0.17), Inches(2.05), Inches(0.42), Inches(0.3), 11, color=color, bold=True)
            add_text(slide, heading, x + Inches(0.17), Inches(2.42), Inches(1.75), Inches(0.3), 15, color=NAVY, bold=True)
            add_text(slide, detail, x + Inches(0.17), Inches(2.82), Inches(1.75), Inches(0.25), 11, color=MUTED)
            if index < len(steps) - 1:
                add_text(slide, "→", x + Inches(2.12), Inches(2.32), Inches(0.38), Inches(0.45), 22, color=LINE, bold=True, align=PP_ALIGN.CENTER)
        add_rect(slide, Inches(0.72), Inches(3.72), Inches(7.15), Inches(2.33), PALE_BLUE, radius=True)
        add_text(slide, "LIMBER PROJECTION", Inches(1.05), Inches(4.05), Inches(2.1), Inches(0.3), 11, color=SKY, bold=True)
        add_text(slide, "Cℓᴬᴮ ≈ ∫ dχ / χ²  Wᴬ(χ) Wᴮ(χ)\n            × P((ℓ + 1/2) / χ, z)", Inches(1.05), Inches(4.55), Inches(6.2), Inches(1.0), 22, color=NAVY, bold=True, font=TITLE_FONT)
        add_image_contain(slide, figures["pk"], Inches(8.18), Inches(3.55), Inches(4.48), Inches(2.75))
        add_text(slide, "ℓ = 10 ... 1,000  |  24 cross-spectra", Inches(8.55), Inches(6.3), Inches(3.7), Inches(0.3), 13, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
        add_footer(slide, 6)
        add_notes(slide, notes[6])

        # Slide 7
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 7, "Why FRB Auto-Correlation Is Not Enough", "Noise floor")
        add_image_contain(slide, figures["shot"], Inches(0.55), Inches(1.67), Inches(8.2), Inches(4.9))
        add_rect(slide, Inches(8.95), Inches(1.92), Inches(3.7), Inches(4.25), NAVY, radius=True)
        add_text(slide, "SHOT NOISE", Inches(9.28), Inches(2.28), Inches(2.9), Inches(0.35), 13, color=GOLD, bold=True)
        add_text(slide, "Nℓ = 1 / n̄", Inches(9.28), Inches(2.92), Inches(2.9), Inches(0.65), 30, color=WHITE, bold=True, font=TITLE_FONT)
        add_text(slide, "~10³", Inches(9.28), Inches(3.95), Inches(2.9), Inches(0.62), 31, color=CORAL, bold=True, font=TITLE_FONT)
        add_text(slide, "above the clustering signal\nin the shallow survey", Inches(9.28), Inches(4.62), Inches(2.8), Inches(0.8), 15, color=RGBColor(219, 228, 232))
        add_text(slide, "Cross-correlation is the core method, not a small optimisation.", Inches(9.28), Inches(5.45), Inches(2.85), Inches(0.55), 12, color=WHITE, bold=True)
        add_footer(slide, 7)
        add_notes(slide, notes[7])

        # Slide 8
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 8, "Tomographic Cross-Signal", "Result I")
        add_image_contain(slide, figures["cross"], Inches(0.55), Inches(1.58), Inches(9.05), Inches(5.35))
        add_rect(slide, Inches(9.78), Inches(1.85), Inches(2.85), Inches(4.78), WHITE, radius=True, line_color=LINE)
        add_text(slide, "24", Inches(10.12), Inches(2.15), Inches(2.1), Inches(0.55), 30, color=TEAL, bold=True, font=TITLE_FONT)
        add_text(slide, "CROSS-SPECTRA", Inches(10.12), Inches(2.76), Inches(2.1), Inches(0.28), 11, color=MUTED, bold=True)
        add_rect(slide, Inches(10.12), Inches(3.35), Inches(1.8), Inches(0.04), LINE)
        add_text(slide, "LOW z", Inches(10.12), Inches(3.68), Inches(1.0), Inches(0.25), 11, color=CORAL, bold=True)
        add_text(slide, "Shallow survey\nhas more signal", Inches(10.12), Inches(4.02), Inches(1.95), Inches(0.62), 15, color=INK)
        add_text(slide, "HIGH z", Inches(10.12), Inches(4.92), Inches(1.0), Inches(0.25), 11, color=TEAL, bold=True)
        add_text(slide, "Deep survey\nretains information", Inches(10.12), Inches(5.26), Inches(1.95), Inches(0.62), 15, color=INK)
        add_footer(slide, 8)
        add_notes(slide, notes[8])

        # Slide 9
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 9, "Fisher Forecast", "Inference")
        add_bullets(slide, ["Parameters θ = (b₀, δ)", "1% central finite differences", "Cosmology held fixed", "Local Gaussian likelihood"], Inches(0.72), Inches(1.95), Inches(4.3), Inches(2.7), size=19)
        add_rect(slide, Inches(0.72), Inches(5.02), Inches(4.2), Inches(1.08), PALE_TEAL, radius=True)
        add_text(slide, "FoM = 1 / √det(Cov)", Inches(1.02), Inches(5.35), Inches(3.6), Inches(0.45), 23, color=TEAL, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
        chart_x, chart_y = Inches(5.65), Inches(1.82)
        add_rect(slide, chart_x, chart_y, Inches(6.7), Inches(4.65), WHITE, radius=True, line_color=LINE)
        add_rect(slide, chart_x + Inches(0.82), chart_y + Inches(3.72), Inches(4.95), Inches(0.025), MUTED)
        add_rect(slide, chart_x + Inches(0.82), chart_y + Inches(0.65), Inches(0.025), Inches(3.1), MUTED)
        ellipse1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, chart_x + Inches(1.65), chart_y + Inches(1.15), Inches(3.45), Inches(2.0))
        ellipse1.fill.background(); ellipse1.line.color.rgb = PALE_CORAL; ellipse1.line.width = Pt(8); ellipse1.rotation = 25
        ellipse2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, chart_x + Inches(2.23), chart_y + Inches(1.53), Inches(2.3), Inches(1.3))
        ellipse2.fill.background(); ellipse2.line.color.rgb = CORAL; ellipse2.line.width = Pt(4); ellipse2.rotation = 25
        add_text(slide, "b₀", chart_x + Inches(5.86), chart_y + Inches(3.55), Inches(0.45), Inches(0.3), 16, color=NAVY, bold=True)
        add_text(slide, "δ", chart_x + Inches(0.38), chart_y + Inches(0.42), Inches(0.3), Inches(0.3), 16, color=NAVY, bold=True)
        add_text(slide, "smaller ellipse  →  stronger joint constraint", chart_x + Inches(1.45), chart_y + Inches(4.02), Inches(4.5), Inches(0.3), 13, color=MUTED, align=PP_ALIGN.CENTER)
        add_footer(slide, 9)
        add_notes(slide, notes[9])

        # Slide 10
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 10, "Cross-Correlation Enables the Measurement", "Main result")
        add_image_contain(slide, figures["fisher"], Inches(0.35), Inches(1.52), Inches(8.3), Inches(5.48))
        add_metric(slide, "17-51×", "SMALLER 1D ERRORS", Inches(8.92), Inches(1.82), Inches(3.15), CORAL)
        add_metric(slide, "61-449×", "LARGER FIGURE OF MERIT", Inches(8.92), Inches(3.08), Inches(3.15), TEAL)
        add_rect(slide, Inches(8.92), Inches(4.47), Inches(3.15), Inches(1.8), NAVY, radius=True)
        add_text(slide, "KiDS result", Inches(9.18), Inches(4.74), Inches(2.6), Inches(0.26), 11, color=GOLD, bold=True)
        add_text(slide, "Bias becomes measurable,\nbut the models still overlap.", Inches(9.18), Inches(5.18), Inches(2.55), Inches(0.75), 17, color=WHITE, bold=True)
        add_text(slide, "Only 3.3% sky overlap", Inches(9.18), Inches(6.0), Inches(2.55), Inches(0.24), 10, color=RGBColor(196, 211, 218))
        add_footer(slide, 10)
        add_notes(slide, notes[10])

        # Slide 11
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 11, "From KiDS to LSST", "Outlook")
        add_image_contain(slide, figures["lsst"], Inches(0.35), Inches(1.55), Inches(8.3), Inches(5.35))
        add_rect(slide, Inches(8.85), Inches(1.82), Inches(3.65), Inches(4.85), WHITE, radius=True, line_color=LINE)
        add_text(slide, "SURVEY SCALE", Inches(9.18), Inches(2.12), Inches(2.8), Inches(0.28), 11, color=MUTED, bold=True)
        add_rich_text(slide, [("1,347", NAVY, True), (" → ", MUTED, False), ("18,000", TEAL, True), (" deg²", MUTED, False)], Inches(9.18), Inches(2.58), Inches(2.9), Inches(0.5), size=23)
        add_rich_text(slide, [("6", NAVY, True), (" → ", MUTED, False), ("10", TEAL, True), (" bins", MUTED, False)], Inches(9.18), Inches(3.15), Inches(2.9), Inches(0.5), size=23)
        add_rect(slide, Inches(9.18), Inches(3.86), Inches(2.65), Inches(0.04), LINE)
        add_text(slide, "DEEP MAGNETAR", Inches(9.18), Inches(4.18), Inches(2.6), Inches(0.28), 11, color=CORAL, bold=True)
        add_text(slide, "σ(b₀)  1.27 → 0.25", Inches(9.18), Inches(4.62), Inches(2.7), Inches(0.34), 17, color=NAVY, bold=True)
        add_text(slide, "FoM     1.59 → 42.01", Inches(9.18), Inches(5.12), Inches(2.7), Inches(0.34), 17, color=NAVY, bold=True)
        add_text(slide, "Models separate at 1σ", Inches(9.18), Inches(5.78), Inches(2.65), Inches(0.35), 14, color=TEAL, bold=True)
        add_footer(slide, 11)
        add_notes(slide, notes[11])

        # Slide 12
        slide = presentation.slides.add_slide(blank)
        set_background(slide)
        add_title(slide, 12, "Limitations", "Interpretation")
        limitations = [
            ("ℓ ≥ 10", "Limber approximation", PALE_BLUE, SKY),
            ("Gaussian", "Optimistic Fisher bounds", PALE_CORAL, CORAL),
            ("n(z), b(z)", "Model assumptions", PALE_TEAL, TEAL),
            ("fsky", "Simplified geometry", RGBColor(240, 234, 214), GOLD),
            ("Ncross ≈ 0", "Cross-shot noise neglected", RGBColor(235, 232, 239), NAVY),
        ]
        positions = [(0.72, 1.92), (4.6, 1.92), (8.48, 1.92), (2.66, 4.28), (6.54, 4.28)]
        for (headline, detail, fill, accent), (x, y) in zip(limitations, positions):
            add_rect(slide, Inches(x), Inches(y), Inches(3.25), Inches(1.72), fill, radius=True)
            add_text(slide, headline, Inches(x + 0.25), Inches(y + 0.25), Inches(2.75), Inches(0.42), 22, color=accent, bold=True, font=TITLE_FONT, align=PP_ALIGN.CENTER)
            add_text(slide, detail, Inches(x + 0.25), Inches(y + 0.94), Inches(2.75), Inches(0.4), 14, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, "Forecasts are informative, but should be read as best-case precision.", Inches(2.3), Inches(6.45), Inches(8.75), Inches(0.4), 16, color=MUTED, align=PP_ALIGN.CENTER)
        add_footer(slide, 12)
        add_notes(slide, notes[12])

        # Slide 13
        slide = presentation.slides.add_slide(blank)
        set_background(slide, NAVY)
        add_text(slide, "CONCLUSION", Inches(0.72), Inches(0.58), Inches(2.3), Inches(0.3), 12, color=GOLD, bold=True)
        add_text(slide, "FRB environments can be studied\nwithout localising every burst.", Inches(0.72), Inches(1.18), Inches(11.7), Inches(1.25), 31, color=WHITE, bold=True, font=TITLE_FONT)
        items = [
            ("01", "Auto-correlation", "Shot-noise limited", CORAL),
            ("02", "KiDS tomography", "FoM improves 61-449×", TEAL),
            ("03", "LSST scale", "Progenitor models separate", GOLD),
        ]
        for index, (number, heading, detail, color) in enumerate(items):
            x = Inches(0.72 + index * 4.15)
            add_rect(slide, x, Inches(3.08), Inches(3.65), Inches(2.05), RGBColor(38, 67, 91), radius=True, line_color=RGBColor(65, 91, 110))
            add_text(slide, number, x + Inches(0.28), Inches(3.36), Inches(0.5), Inches(0.35), 12, color=color, bold=True)
            add_text(slide, heading, x + Inches(0.28), Inches(3.88), Inches(3.0), Inches(0.36), 18, color=WHITE, bold=True)
            add_text(slide, detail, x + Inches(0.28), Inches(4.42), Inches(3.0), Inches(0.35), 14, color=RGBColor(204, 217, 223))
        add_rect(slide, Inches(0.72), Inches(5.82), Inches(11.95), Inches(0.06), TEAL)
        add_text(slide, "FRBs  +  galaxy tomography  →  progenitor physics", Inches(0.72), Inches(6.18), Inches(9.6), Inches(0.48), 20, color=WHITE, bold=True)
        add_text(slide, "Thank you", Inches(10.62), Inches(6.18), Inches(2.0), Inches(0.42), 17, color=GOLD, bold=True, align=PP_ALIGN.RIGHT)
        add_notes(slide, notes[13])

        presentation.core_properties.title = "Probing Fast Radio Burst Environments"
        presentation.core_properties.subject = "Bachelor thesis presentation"
        presentation.core_properties.author = "Will Ruppert"
        presentation.core_properties.keywords = "FRB, KiDS, cross-correlation, Fisher forecast"
        presentation.save(OUTPUT_PATH)

    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    build_presentation()